# Copyright (c) 2026, Nirmaan (Stratos Infra Technologies Pvt. Ltd.) and contributors
# For license information, please see license.txt

"""SLICE 1e -- the Excel half of the rate-file round trip. PURE: bytes in, rows out; rows in, bytes out.

WHY EXCEL BY DEFAULT (owner X-a, 2026-09-22). The owner downloaded the HVAC CSV, added one row in Excel
and uploaded it: Excel had rewritten the panel ratios "1:6" and "1:4" as the TIMES 01:06 and 01:04 on
the way through, and the spec reader (correctly) refused them. A CSV carries no cell types, so Excel
guesses -- and a guess it makes silently is exactly what a rate file cannot afford. An .xlsx carries
the type: every text column here is written with the TEXT number format ("@"), at the CELL and at the
COLUMN, so a value is shown and re-read as typed, and a row a user adds below the data inherits the
column's text format before they type into it.

WHAT IS TEXT AND WHAT IS A NUMBER is decided by the CALLER (the exporter knows the discipline's column
spaces): the identity columns, the category / kind labels, `item_name` / `item_detail` and every
attribute whose definition is not numeric are TEXT; the rate and markup columns and the numeric
attributes are NUMBERS, written exactly (a float is a float, never rounded or formatted).

READING is the mirror and is deliberately DUMB: a cell comes back as the string the CSV path would
have read -- text as is, a number as `str()` (openpyxl hands back `int` for an integral float, so a
stored `2.0` comes back as "2", which `csv_importer.coerce_rate` reads as 2.0 -- the same value), a
date or time as its ISO text so it SURFACES as a change or a refusal rather than vanishing. Nothing is
repaired here; the importer's preview is the defence, as it always was.

DETECTION IS BY CONTENT, never by file name: an .xlsx is a zip and starts with the four bytes
`PK\\x03\\x04`; anything else is handed to the CSV decoder unchanged.

openpyxl is already a dependency of frappe and of this app (VERIFIED: 3.1.5 in the bench env,
`Required-by: frappe, nirmaan_stack`). No new dependency.
"""

import datetime
import io

import openpyxl
from openpyxl.utils import get_column_letter

XLSX_MAGIC = b"PK\x03\x04"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CONTENT_TYPE = "text/csv"
TEXT_FORMAT = "@"
SHEET_TITLE = "rate_master"


def is_xlsx(raw):
    """True when the bytes are a zip container (every .xlsx is one). A str is never an xlsx."""
    return isinstance(raw, (bytes, bytearray)) and bytes(raw[:4]) == XLSX_MAGIC


def write_xlsx(headers, rows, numeric_columns):
    """bytes of a one-sheet workbook. `rows` are lists of raw values aligned with `headers`; a column
    named in `numeric_columns` is written as a NUMBER when its value is an int / float (a None is an
    empty cell; anything else falls back to text); every other column is TEXT, cell and column."""
    numeric = set(numeric_columns or ())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ws.append(list(headers))
    for col_idx, name in enumerate(headers, start=1):
        if name not in numeric:
            ws.column_dimensions[get_column_letter(col_idx)].number_format = TEXT_FORMAT
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, (name, value) in enumerate(zip(headers, row), start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if value is None or value == "":
                if name not in numeric:
                    cell.number_format = TEXT_FORMAT
                continue
            if name in numeric and isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.value = value
            else:
                cell.value = str(value)
                cell.number_format = TEXT_FORMAT
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(value):
    """The string the CSV path would have read for this cell. Never a repair: a datetime becomes its
    ISO text so the importer SEES it (as a change, or a refusal by the spec reader)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def read_xlsx(raw):
    """(headers, data_rows) in exactly the shape `csv_importer.parse_csv_text` returns: headers
    stripped, data_rows = [(1-based data row number, [cell text, ...]), ...]. The FIRST worksheet is
    read; a workbook with no rows yields ([], [])."""
    wb = openpyxl.load_workbook(io.BytesIO(bytes(raw)), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    # openpyxl pads to the sheet's max column; trailing fully-empty rows are dropped, blank cells kept
    while rows and not any(c not in (None, "") for c in rows[-1]):
        rows.pop()
    if not rows:
        return [], []
    headers = [(_cell_text(h)).strip() for h in rows[0]]
    # a trailing run of blank HEADER cells is not a column (Excel widens the used range easily)
    while headers and headers[-1] == "":
        headers.pop()
    width = len(headers)
    data = []
    for i, r in enumerate(rows[1:], start=1):
        cells = [_cell_text(c) for c in r[:width]]
        if len(r) > width and any(c not in (None, "") for c in r[width:]):
            # a value beyond the header's width is kept so the importer reports it as it does for CSV
            cells = [_cell_text(c) for c in r]
        data.append((i, cells))
    return headers, data
