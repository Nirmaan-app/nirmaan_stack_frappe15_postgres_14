#!/usr/bin/env python3
"""SLICE 12d-3 -- THE INSULATION AUDIT, THE ANALYSIS + THE EXCEL (host; pure; no AI, no DB).

Inputs: the rows `audit12d3_build_rows.py` captured, the two-path pricing `audit12d3_pricing.ts` produced, the
asset the live config equals (stocked sizes per family), the 12d-2 sample fixture (the 25-row repeatability
check) and the cost summary. Outputs: one Excel line per row (every capture column, filterable), an analysis
JSON (every automatic check by row), a markdown digest of the automatic findings, and the stratified
hand-review dump (60 rows) for a person to read.

Usage: python audit12d3_analyse.py <audit_rows.json> <audit_pricing.json> <asset.json> <sample12d2.json> <cost.json> <out_dir>
"""
import collections
import json
import random
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

rows_path, pricing_path, asset_path, sample_path, cost_path, out_dir = sys.argv[1:7]
# optional 7th arg: the hand review -- per-row verdicts (item / cladding / thickness / pipe_size) and the
# second-opinion judgments (SO_RIGHT / MODEL_RIGHT / ARGUABLE); merged into the Excel and the digest
HAND = json.load(open(sys.argv[7], encoding="utf-8")) if len(sys.argv) > 7 else {"rows": {}, "so": {}}
ROWS = json.load(open(rows_path, encoding="utf-8"))
PRICING = {p["id"]: p for p in json.load(open(pricing_path, encoding="utf-8"))}
ASSET = json.load(open(asset_path, encoding="utf-8"))
SAMPLE = {(r["boq"], r["sheet"], int(r["excel_row"])): r for r in json.load(open(sample_path, encoding="utf-8"))}
COST = json.load(open(cost_path, encoding="utf-8"))

CFG = next(c for c in ASSET["category_configs"] if c["category_id"] == "hvac_insulation")
PIPE_FAMILIES = {"Nitrile Rubber Insulation", "Tubular Puf Insulation", "Cladding Only"}
STOCK = collections.defaultdict(lambda: {"th": set(), "pipe": set(), "clad": set()})
for it in ASSET["items"]:
    if it["kind"] != "hvac_insulation_item":
        continue
    a = it["attributes"]
    f = STOCK[a.get("item")]
    if a.get("thickness_mm") is not None:
        f["th"].add(float(a["thickness_mm"]))
    if a.get("pipe_size_mm") is not None:
        f["pipe"].add(float(a["pipe_size_mm"]))
    if a.get("cladding"):
        f["clad"].add(a["cladding"])


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def note_texts(block):
    out = []
    if isinstance(block, dict):
        for k, v in block.items():
            if isinstance(v, dict):
                out += [str(x) for x in v.values() if x]
            elif isinstance(v, list):
                out += [str(x) for x in v if x]
            elif v:
                out.append(str(v))
    elif isinstance(block, list):
        out += [str(x) for x in block if x]
    elif block:
        out.append(str(block))
    return out


def own_text(r):
    p = r.get("payload") or {}
    return norm(" | ".join([r.get("description") or ""] + note_texts(p.get("notes"))))


def heading_text(r):
    p = r.get("payload") or {}
    parts = []
    for a in p.get("ancestor_chain") or []:
        if a.get("relation") == "sheet":
            continue
        parts.append(a.get("description") or "")
        parts += note_texts(a.get("notes"))
    return norm(" | ".join(parts))


def model_attrs(r):
    items = ((r.get("answer") or {}).get("items")) or []
    out = []
    for it in items:
        out.append({k: (v or {}).get("value") for k, v in (it.get("attributes") or {}).items()})
    return out


def state_of(v):
    if v is None:
        return "absent"
    if v == "None":
        return "not mentioned"
    return "stated"


NUM = r"(\d+(?:\.\d+)?)"
TH_OWN = re.compile(NUM + r"\s*(?:mm|m\.m\.?)\.?\s*(?:thk\.?|thick|thickness|th\b)", re.I)
TH_OWN2 = re.compile(r"(?:thickness|thick|thk)\.?\s*(?:of\s*)?" + NUM + r"\s*mm", re.I)
CLAD_WORDS = re.compile(r"alumin[iu]um|\balu\b|\b2[246]\s*g\b|\b2[246]\s*gauge|foil|glass\s*cloth|gi\s*frame|g\.i\.?\s*frame|perforated|cladd|\bclad\b|canvas|tissue|uv\s*shield|coating|gss", re.I)
PIPE_WORDS = re.compile(NUM + r"\s*(?:mm\s*)?(?:nb\b|dia\b|dia\.|od\b|\"|inch|”)", re.I)
MAT = [
    (re.compile(r"xlpe|\beps\b|rock\s*wool|rockwool|mineral\s*wool|phenolic|\bpir\b|polyethylene|epdm|thermocol|expanded\s*polystyrene", re.I), "none of these"),
    (re.compile(r"fib(?:er|re)\s*glass|fiberglass|fibreglass|glass\s*wool|glasswool", re.I), "fiberglass"),
    (re.compile(r"\bpuf\b|polyurethane", re.I), "puf"),
    (re.compile(r"nitril|elastomeric|arma\s*-?\s*flex|aeroflex|k-?\s*flex|insulflex|superlon|closed\s*cell|open\s*cell", re.I), "nitrile"),
]
SCHEDULE = re.compile(NUM + r"\s*(?:mm|nb)?\s*(?:to|-|–|upto|up to)\s*" + NUM + r"\s*(?:mm|nb)?\s*:", re.I)
SLASH_LIST = re.compile(r"^\s*\d+(?:\.\d+)?\s*(?:/\s*\d+(?:\.\d+)?\s*)+(?:mm)?", re.I)
LAYERS = re.compile(r"\+|\bx\s*\d|\d\s*x\b|layer|double", re.I)


def numbers_in(s):
    return [float(x) for x in re.findall(NUM, str(s or ""))]


def expected_family(text, unit_class):
    for rx, kind in MAT:
        if rx.search(text):
            if kind == "none of these":
                return {"none of these"}
            if kind == "fiberglass":
                return {" Fiberglass Rigid Board Insulation, Density 48Kg/m3"}
            if kind == "puf":
                return {"Tubular Puf Insulation"}
            if kind == "nitrile":
                if unit_class == "length":
                    return {"Nitrile Rubber Insulation"}
                return {"Thermal Nitrile Insulation", "Acoustic Nitrile Insulation"}
    return None


def reason_class(reason):
    if not reason:
        return ""
    s = re.sub(r"\d+(?:\.\d+)?", "N", reason)
    s = re.sub(r"'[^']*'", "'X'", s)
    s = re.sub(r"for [A-Za-z0-9 ,.\-]+ - price this row by hand", "for X - price this row by hand", s)
    return s[:120]


analysis = []
for r in ROWS:
    pid = r["id"]
    pr = PRICING.get(pid) or {}
    panel = pr.get("panel") or {}
    calc = pr.get("calculator") or {}
    own = own_text(r)
    head = heading_text(r)
    ma = model_attrs(r)
    m0 = ma[0] if ma else {}
    unit_class = panel.get("unitClass")
    a = {"id": pid, "boq": r["boq"], "sheet": r["sheet"], "excel_row": r["excel_row"], "unit": r.get("unit"),
         "n_model_items": len(ma), "checks": {}, "rules_fired": [], "violations": [], "reading_failures": []}
    ck = a["checks"]

    # ── 2(a) THE MODEL'S READING, automatic ──────────────────────────────────────────────────────
    own_th = sorted(set(numbers_in(" ".join(TH_OWN.findall(own) + TH_OWN2.findall(own)))))
    model_th = [x.get("thickness_mm") for x in ma]
    model_th_nums = sorted(set(n for v in model_th for n in numbers_in(v)))
    if own_th:
        if any(v is None for v in model_th) and not model_th_nums:
            ck["S3_own_thickness"] = "NOT USED: model left thickness out"
            a["reading_failures"].append(f"thickness: row states {own_th}, model left it out")
        elif set(own_th) & set(model_th_nums):
            ck["S3_own_thickness"] = "used" if len(own_th) == 1 else f"used (one of several own mentions {own_th}; hand check)"
        else:
            ck["S3_own_thickness"] = f"NOT USED: row states {own_th}, model wrote {model_th}"
            a["reading_failures"].append(ck["S3_own_thickness"])
    else:
        ck["S3_own_thickness"] = "n/a (row states none)"
    # material
    exp = expected_family(own, unit_class)
    fams = [x.get("item") for x in ma]
    if exp is not None:
        if any(f in exp for f in fams):
            ck["material_own"] = "carried"
        elif all(f is None for f in fams):
            ck["material_own"] = f"NOT CARRIED: row names a material ({sorted(exp)}), model left the family out"
            a["reading_failures"].append(ck["material_own"])
        else:
            ck["material_own"] = f"MISMATCH?: row suggests {sorted(exp)}, model {fams}"
            a["reading_failures"].append(ck["material_own"])
    else:
        ck["material_own"] = "n/a (no material word in the row)"
    # cladding
    clad_vals = [x.get("cladding") for x in ma]
    own_clad = bool(CLAD_WORDS.search(own))
    head_clad = bool(CLAD_WORDS.search(head))
    if own_clad and all(v == "None" for v in clad_vals):
        ck["cladding"] = "NOT-MENTIONED CONTRADICTED: the row's own text has a cladding word"
        a["reading_failures"].append(ck["cladding"] + f" ({CLAD_WORDS.search(own).group(0)!r})")
    elif own_clad and all(v is None for v in clad_vals):
        ck["cladding"] = "left out while the row has a cladding word (R3 could-not-map is legitimate; hand check)"
    elif own_clad:
        ck["cladding"] = "carried"
    elif head_clad and all(v == "None" for v in clad_vals):
        ck["cladding"] = "'None' while a HEADING has a cladding word (hand check: does the heading describe this row?)"
    elif not own_clad and not head_clad and any(v not in (None, "None", "No") for v in clad_vals):
        ck["cladding"] = f"VALUE WITHOUT A MENTION?: model {clad_vals} but no cladding word in row or headings"
        a["reading_failures"].append(ck["cladding"])
    elif not own_clad and not head_clad and any(v == "No" for v in clad_vals):
        ck["cladding"] = "model 'No' with no cladding word in row or headings (hand check: 'No' should mean the row STATES none; the price is the same as the R1 default)"
    else:
        ck["cladding"] = "n/a or consistent"
    # pipe size
    pipe_vals = [x.get("pipe_size_mm") for x in ma]
    own_pipe = PIPE_WORDS.search(own)
    if unit_class == "length" and own_pipe and all(v is None for v in pipe_vals):
        ck["pipe_size"] = f"NOT CARRIED: row states {own_pipe.group(0)!r}, model left pipe size out"
        a["reading_failures"].append(ck["pipe_size"])
    elif unit_class == "length" and own_pipe:
        ck["pipe_size"] = "carried"
    elif unit_class == "length" and not own_pipe and any(v for v in pipe_vals):
        ck["pipe_size"] = "from a heading (row states none)" if PIPE_WORDS.search(head) else "VALUE WITHOUT A MENTION?"
        if ck["pipe_size"] == "VALUE WITHOUT A MENTION?":
            a["reading_failures"].append("pipe size: " + ck["pipe_size"])
    else:
        ck["pipe_size"] = "n/a"
    # three states per attribute (first item)
    a["states"] = {k: state_of(m0.get(k)) for k in ("item", "cladding", "thickness_mm", "pipe_size_mm", "brand", "material_as_written")}
    a["schedule_heading"] = bool(SCHEDULE.search(head))
    a["several_values"] = any(SLASH_LIST.match(str(v or "")) for v in model_th)
    a["double_layer_text"] = any(LAYERS.search(str(v or "")) for v in model_th)
    a["foil_mentioned"] = bool(re.search(r"foil", own + " " + head))
    a["so_verdict"] = r.get("so_verdict")
    a["so_flags"] = [f for f in (r.get("item_flags") or []) if f.get("check") == "second_opinion"]
    a["text_flags"] = r.get("text_flags") or []

    # ── 3 THE RULES, automatic (from the panel path) ─────────────────────────────────────────────
    blocks = panel.get("items") or []
    a["panel_priced"] = bool(panel.get("rowPriced"))
    a["panel_reason"] = panel.get("reason") or (blocks[0].get("reason") if blocks and blocks[0].get("reason") else None)
    a["panel_totals"] = panel.get("totals")
    a["calc_priced"] = bool(calc.get("rowPriced"))
    a["calc_totals"] = calc.get("totals")
    a["calc_reason"] = calc.get("reason") or ((calc.get("items") or [{}])[0].get("reason") if calc.get("items") else None)
    a["families_priced"] = [b.get("family") for b in blocks]
    a["skus"] = [b.get("skuLine") for b in blocks]
    a["working"] = [b.get("working") for b in blocks]
    a["divergences"] = pr.get("divergences") or []
    a["divergence_cause"] = pr.get("divergence_cause")
    if a["divergences"]:
        pr_ = (a["panel_reason"] or "")
        cr_ = (a["calc_reason"] or "")
        if "no SKU for this combination" in pr_ and cr_.startswith("could not tell"):
            a["divergence_cause"] = "C_option_not_offered (the row-290 class, ACCEPTED by owner F2)"
        elif "several values stated" in cr_ and "+" in cr_:
            a["divergence_cause"] = "T6_typed_layers (RULED 12d-1b: the calculator's typed layers stay refused; the panel reads model cells)"
    if panel.get("unitNote"):
        a["rules_fired"].append("unit rule: " + panel["unitNote"])
    for bi, b in enumerate(blocks):
        mi = ma[bi] if bi < len(ma) else {}
        if b.get("familyDefaulted"):
            a["rules_fired"].append(f"family default -> {b['familyDefaulted']['value']} ({b['familyDefaulted']['rule']})")
            if mi.get("item") not in (None, "None"):
                a["violations"].append(f"block {bi}: family default fired over a STATED family {mi.get('item')!r}")
        for f in b.get("fields") or []:
            fid = f["id"]
            mv = mi.get(fid) if mi else None
            if f.get("defaulted"):
                a["rules_fired"].append(f"default {fid} -> {f['value']} ({f.get('rule')})")
                if mv not in (None, "None"):
                    a["violations"].append(f"block {bi}: default on {fid} fired over a STATED value {mv!r}")
                if fid == "cladding" and mv is None:
                    a["violations"].append(f"block {bi}: cladding default fired over an ABSENT answer (absent_as_none is false)")
            if f.get("note"):
                a["rules_fired"].append(f"{fid}: {f['note']}")
            fam = b.get("family")
            stock = STOCK.get(fam) if fam else None
            if stock and f.get("value") not in (None, "", "None") and fid in ("thickness_mm", "pipe_size_mm") and b.get("state") == "priced":
                try:
                    val = float(f["value"])
                except ValueError:
                    val = None
                key = "th" if fid == "thickness_mm" else "pipe"
                comp = re.search(r"priced as ([\d.\s+]+) mm", f.get("note") or "")
                if comp and "+" not in comp.group(1):
                    comp = None   # a single 'priced as N mm' is a ladder hop, not a composition
                if comp and fid == "thickness_mm":
                    layers = [float(x) for x in re.findall(NUM, comp.group(1))]
                    stated = numbers_in(mv)
                    tot = sum(layers)
                    if len(layers) > 4:
                        a["violations"].append(f"composition > 4 layers: {layers}")
                    if any(l not in stock["th"] for l in layers):
                        a["violations"].append(f"composition layer not stocked: {layers} vs {sorted(stock['th'])}")
                    if stated and abs(tot - max(stated)) > 2.0 + 1e-9 and "layers" not in (f.get("note") or ""):
                        a["violations"].append(f"composition {layers} = {tot} is more than 2 mm from the stated {stated}")
                    a["rules_fired"].append(f"composition: {layers} (stated {stated})")
                elif val is not None and stock[key] and val not in stock[key]:
                    a["violations"].append(f"{fid} resolved to {val}, NOT a stocked {key} size of {fam}: {sorted(stock[key])}")
            if fid == "cladding" and mi:
                if mv == "Aluminium Foil" and f.get("value") == "26G Aluminium":
                    a["rules_fired"].append("R4 foil on a pipe -> 26G")
                    if fam not in PIPE_FAMILIES:
                        a["violations"].append(f"R4 foil->26G applied on a non-pipe family {fam!r}")
                if mv == "Aluminium Foil" and fam == "Acoustic Nitrile Insulation" and "foil on an acoustic row" not in (b.get("reason") or ""):
                    a["violations"].append("acoustic + foil did NOT refuse by R4")
            if fid == "thickness_mm" and mv and SLASH_LIST.match(str(mv)):
                # T2: a bare slash list prices from its HIGHEST value (the reader), then the ladder / composition
                hi = max(numbers_in(mv))
                said = numbers_in((f.get("note") or "") + " " + " ".join(w for w in (b.get("working") or []) if w.startswith("BoQ says")))
                if b.get("state") == "priced" and (hi in said or (f.get("value") not in (None, "") and float(f["value"]) >= hi)):
                    a["rules_fired"].append(f"T2 several values {mv!r} -> the highest ({hi:g})")
                elif b.get("state") == "priced":
                    a["violations"].append(f"several values {mv!r} priced from a value below the highest {hi:g} (value {f.get('value')})")
        if b.get("state") == "blank" and b.get("reason"):
            a["rules_fired"].append("refusal: " + b["reason"])
        for w in b.get("working") or []:
            if w.startswith("BoQ says") or "layers" in w or "R4" in w:
                a["rules_fired"].append("working: " + w)
    # outer-only cladding on a composition (the inner layers carry cladding No)
    if len(blocks) > 1 and len(ma) == 1:
        inner = [b for b in blocks[:-1]]
        for b in inner:
            cf = next((f for f in b.get("fields") or [] if f["id"] == "cladding"), None)
            if cf and cf.get("value") not in ("No", None, ""):
                a["violations"].append(f"inner layer carries cladding {cf.get('value')!r} (outer_only)")
    # unit rule
    u = norm(r.get("unit"))
    if (not u or re.sub(r"[^a-z0-9]", "", u) in ("ro", "rateonly")) and not panel.get("unitNote") and panel.get("declined") is False:
        if not (a["panel_reason"] or "").startswith("no unit") and "unit" not in (a["panel_reason"] or ""):
            a["violations"].append("row with no / rate-only unit carries no unit note")
    # hand-typed rates
    hand = r.get("hand_rates") or []
    hs = {}
    for c in hand:
        k = norm(c.get("kind"))
        if "supply" in k:
            hs["supply"] = c.get("rate")
        elif "install" in k:
            hs["install"] = c.get("rate")
        elif "combined" in k or "total" in k:
            hs["combined"] = c.get("rate")
    a["hand"] = hs
    t = a["panel_totals"] or {}
    a["hand_diff"] = {k: (t.get(k + "_rate") - hs[k]) for k in hs if t.get(k + "_rate") is not None and isinstance(hs[k], (int, float))}
    # repeatability (the 25 sampled rows)
    s = SAMPLE.get((r["boq"], r["sheet"], int(r["excel_row"])))
    if s:
        old = [{k: (v or {}).get("value") for k, v in (it.get("attributes") or {}).items()} for it in (s.get("answer") or {}).get("items") or []]
        diffs = []
        for i in range(max(len(old), len(ma))):
            o = old[i] if i < len(old) else {}
            n = ma[i] if i < len(ma) else {}
            for k in sorted(set(o) | set(n)):
                if norm(o.get(k)) != norm(n.get(k)):
                    diffs.append(f"item{i}.{k}: 12d-2 {o.get(k)!r} -> 12d-3 {n.get(k)!r}")
        a["repeat_diffs"] = diffs
        a["repeat_so_12d2"] = [f for f in (s.get("item_flags") or []) if f.get("check") == "second_opinion"]
    analysis.append(a)

# ── summaries ────────────────────────────────────────────────────────────────────────────────────
def pct(n, d):
    return f"{(100.0 * n / d):.1f}%" if d else "-"


total = len(analysis)
priced = sum(1 for a in analysis if a["panel_priced"])
refused = [a for a in analysis if not a["panel_priced"]]
by_reason = collections.Counter(reason_class(a["panel_reason"]) for a in refused)
by_family = collections.Counter()
by_family_priced = collections.Counter()
for a in analysis:
    fam = a["families_priced"][0] if a["families_priced"] else "(none)"
    by_family[fam] += 1
    if a["panel_priced"]:
        by_family_priced[fam] += 1
by_boq = collections.defaultdict(lambda: [0, 0])
for a in analysis:
    by_boq[a["boq"]][0] += 1
    by_boq[a["boq"]][1] += int(a["panel_priced"])
reading_fail = [a for a in analysis if a["reading_failures"]]
violations = [a for a in analysis if a["violations"]]
diverged = [a for a in analysis if a["divergences"]]
so_dis = [a for a in analysis if a["so_verdict"] == "disagree"]
repeat = [a for a in analysis if "repeat_diffs" in a]
repeat_changed = [a for a in repeat if a["repeat_diffs"]]
states = collections.Counter()
for a in analysis:
    for k, v in a["states"].items():
        states[(k, v)] += 1

md = []
md.append(f"# 12d-3 automatic analysis -- {total} rows\n")
md.append(f"## 1 COVERAGE\npriced {priced} ({pct(priced, total)}), refused {len(refused)}\n")
md.append("refusals by reason:\n" + "\n".join(f"- {n} x {k}" for k, n in by_reason.most_common()))
md.append("\nby family (priced / rows):\n" + "\n".join(f"- {k!r}: {by_family_priced[k]} / {n}" for k, n in by_family.most_common()))
md.append("\nby BoQ (priced / rows):\n" + "\n".join(f"- {k}: {v[1]} / {v[0]}" for k, v in sorted(by_boq.items())))
md.append("\n## 2(a) THE MODEL'S READING -- automatic\n")
md.append("three states (first item): " + ", ".join(f"{k}={v}: {n}" for (k, v), n in sorted(states.items())))
md.append(f"\nrows with at least one automatic reading failure: {len(reading_fail)} ({pct(len(reading_fail), total)})\n")
cnt = collections.Counter()
for a in analysis:
    for k in ("S3_own_thickness", "material_own", "cladding", "pipe_size"):
        cnt[(k, a["checks"][k].split(":")[0])] += 1
md.append("\n".join(f"- {k}: {v} = {n}" for (k, v), n in sorted(cnt.items())))
md.append("\nevery reading failure by row:\n" + "\n".join(f"- {a['id']}: {'; '.join(a['reading_failures'])}" for a in reading_fail))
md.append(f"\n## 2(c) REPEATABILITY -- {len(repeat)} of the 25 sampled rows re-run; {len(repeat_changed)} changed\n")
md.append("\n".join(f"- {a['id']}: {'; '.join(a['repeat_diffs'])}" for a in repeat_changed) or "- (no attribute differs)")
md.append(f"\n## 3 THE RULES -- automatic: {len(violations)} rows with a violation\n")
md.append("\n".join(f"- {a['id']}: {'; '.join(a['violations'])}" for a in violations) or "- none")
rules_cnt = collections.Counter()
for a in analysis:
    for rf in a["rules_fired"]:
        rules_cnt[re.sub(r"\d+(?:\.\d+)?", "N", rf.split("(")[0])[:80]] += 1
md.append("\nrules fired (normalised, count):\n" + "\n".join(f"- {n} x {k}" for k, n in rules_cnt.most_common(40)))
md.append(f"\n## 4 THE SECOND OPINION -- {len(so_dis)} disagree / {sum(1 for a in analysis if a['so_verdict']=='agree')} agree / {sum(1 for a in analysis if a['so_verdict'] is None)} no verdict\n")
md.append("\n".join(f"- {a['id']} [{'priced' if a['panel_priced'] else 'refused'}]: " + " || ".join(f"{f.get('attr')}: {f.get('reason')}" for f in a["so_flags"]) for a in so_dis))
md.append(f"\n## 5 CALCULATOR = PANEL -- {len(diverged)} rows diverge\n")
dc = collections.Counter(a["divergence_cause"] for a in diverged)
md.append("\n".join(f"- {n} x {k}" for k, n in dc.most_common()))
md.append("\n" + "\n".join(f"- {a['id']}: cause {a['divergence_cause']}; " + " || ".join(f"{d['what']}: P[{d['panel'][:160]}] C[{d['calculator'][:160]}]" for d in a["divergences"]) for a in diverged))
md.append("\n## 6 HAND-PRICED (indicative)\n")
hd = [(a, k, v) for a in analysis for k, v in a["hand_diff"].items()]
md.append(f"rows with a hand rate: {sum(1 for a in analysis if a['hand'])}; priced rows with a comparable hand rate: {len(set(a['id'] for a, _, _ in hd))}")
hd.sort(key=lambda x: -abs(x[2]))
md.append("\n".join(f"- {a['id']} {k}: panel {a['panel_totals'].get(k + '_rate')} vs hand {a['hand'][k]} (diff {v:+.0f}); fam {a['families_priced']}; rules {a['rules_fired'][:3]}" for a, k, v in hd[:25]))
if HAND["rows"]:
    md.append(f"\n## 2(b) HAND REVIEW -- {len(HAND['rows'])} rows, per attribute\n")
    for attr in ("item", "cladding", "thickness", "pipe_size"):
        c = collections.Counter(v[attr] for v in HAND["rows"].values())
        md.append(f"- {attr}: " + ", ".join(f"{k} {n}" for k, n in sorted(c.items()))
                  + f" (error rate WRONG {pct(c.get('WRONG', 0), len(HAND['rows']))}, WRONG+ARGUABLE {pct(c.get('WRONG', 0) + c.get('ARGUABLE', 0), len(HAND['rows']))})")
    md.append("\n" + "\n".join(f"- {k}: item {v['item']}, cladding {v['cladding']}, thickness {v['thickness']}, pipe {v['pipe_size']} -- {v['note']}" for k, v in HAND["rows"].items()))
if HAND["so"]:
    c = collections.Counter(v["verdict"] for v in HAND["so"].values())
    priced_ids = {a["id"] for a in analysis if a["panel_priced"]}
    md.append(f"\n## 4(b) SECOND OPINION JUDGED -- {dict(c)}; SO_RIGHT on PRICED rows: {sum(1 for k, v in HAND['so'].items() if v['verdict'] == 'SO_RIGHT' and k in priced_ids)}\n")
    md.append("\n".join(f"- {k} [{v['verdict']}]: {v['note']}" for k, v in HAND["so"].items()))
md.append("\n## 8 COST\n```\n" + json.dumps(COST, indent=1) + "\n```")
open(f"{out_dir}/audit_analysis.md", "w", encoding="utf-8", newline="\n").write("\n".join(md) + "\n")
json.dump(analysis, open(f"{out_dir}/audit_analysis.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False, default=str)

# ── hand-review selection: 60 rows, stratified and named ────────────────────────────────────────
random.seed(12)
sel = collections.OrderedDict()


HAND_REVIEW_N = 60


def take(name, pred, n):
    cands = [a for a in analysis if pred(a) and a["id"] not in sel]
    random.shuffle(cands)
    for a in cands[:n]:
        if len(sel) >= HAND_REVIEW_N:
            return
        sel[a["id"]] = name


for fam in by_family:
    take(f"family {fam!r}", lambda a, f=fam: (a["families_priced"][0] if a["families_priced"] else "(none)") == f, 3)
take("default fired (cladding)", lambda a: any(x.startswith("default cladding") for x in a["rules_fired"]), 3)
take("default fired (thickness)", lambda a: any(x.startswith("default thickness") for x in a["rules_fired"]), 3)
take("default fired (family)", lambda a: any(x.startswith("family default") for x in a["rules_fired"]), 3)
take("schedule heading", lambda a: a["schedule_heading"], 4)
take("several values", lambda a: a["several_values"], 3)
take("double layer", lambda a: a["double_layer_text"], 3)
take("foil", lambda a: a["foil_mentioned"], 4)
for rc in by_reason:
    take(f"refusal: {rc}", lambda a, c=rc: not a["panel_priced"] and reason_class(a["panel_reason"]) == c, 1)
take("second-opinion disagree", lambda a: a["so_verdict"] == "disagree", 8)
take("automatic reading failure", lambda a: bool(a["reading_failures"]), 6)
take("fill (priced)", lambda a: a["panel_priced"], max(0, HAND_REVIEW_N - len(sel)))
by_id = {r["id"]: r for r in ROWS}
dump = []
for pid, why in sel.items():
    r = by_id[pid]
    a = next(x for x in analysis if x["id"] == pid)
    dump.append(f"=== {pid} | unit {r.get('unit')!r} | STRATUM: {why}\nROW: {r['description']}\nOWN NOTES: {note_texts((r.get('payload') or {}).get('notes'))}\nHEADINGS: {r['headings']}\n"
                f"MODEL: {json.dumps(model_attrs(r), ensure_ascii=False)}\nRAW: {json.dumps(r.get('raw_answer'), ensure_ascii=False)[:600]}\nSO: {a['so_verdict']} {json.dumps(a['so_flags'], ensure_ascii=False)}\nTEXT FLAGS: {a['text_flags']}\n"
                f"PANEL: priced={a['panel_priced']} totals={a['panel_totals']} reason={a['panel_reason']} fams={a['families_priced']} skus={a['skus']}\nRULES: {a['rules_fired']}\nVIOLATIONS: {a['violations']}\nAUTO CHECKS: {a['checks']}\n")
open(f"{out_dir}/hand_review_dump.txt", "w", encoding="utf-8", newline="\n").write("\n".join(dump))
json.dump([{"id": k, "stratum": v} for k, v in sel.items()], open(f"{out_dir}/hand_review_selection.json", "w"), indent=1)

# ── THE EXCEL: one line per row, every capture column ───────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "rows"
cols = ["BoQ", "Sheet", "Row", "Unit", "Button reachable", "Run id", "Run status", "Run active", "Description", "Own notes", "Headings (root-first)",
        "Model items (n)", "Model item (family)", "Model cladding", "Model thickness", "Model pipe size", "Model brand", "Model material as written",
        "State: item", "State: cladding", "State: thickness", "State: pipe size", "Raw reply element", "Text-check flags",
        "Second opinion", "Second-opinion flags", "SO failed",
        "Final family (priced)", "Final cladding", "Final thickness", "Final pipe size", "Rules fired", "SKU(s) / layers", "Working",
        "Panel supply", "Panel install", "Panel combined", "Panel refusal", "Calc supply", "Calc install", "Calc combined", "Calc refusal",
        "Divergence", "Divergence cause", "Hand supply", "Hand install", "Diff supply", "Diff install",
        "Check S3 own thickness", "Check material", "Check cladding", "Check pipe size", "Rule violations", "Repeat diffs (25 rows)",
        "Schedule heading", "Several values", "Double layer", "Foil mentioned", "Batch ts", "Batch rows", "Batch in/out tokens", "Failed attempts", "Hand-review stratum",
        "Hand: item", "Hand: cladding", "Hand: thickness", "Hand: pipe size", "Hand: note", "SO judged", "SO judgement note"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDDDDD")
    c.alignment = Alignment(wrap_text=True, vertical="top")
amber = PatternFill("solid", fgColor="FFE699")
red = PatternFill("solid", fgColor="F8CBAD")
for r in ROWS:
    a = next(x for x in analysis if x["id"] == r["id"])
    pr = PRICING.get(r["id"]) or {}
    panel = pr.get("panel") or {}
    blocks = panel.get("items") or []
    ma = model_attrs(r)
    m0 = ma[0] if ma else {}
    def fld(bi, fid):
        if bi >= len(blocks):
            return ""
        f = next((f for f in blocks[bi].get("fields") or [] if f["id"] == fid), None)
        return "" if not f else (f["value"] + (" [default]" if f.get("defaulted") else "") + ((" {" + f["note"] + "}") if f.get("note") else ""))
    t = a["panel_totals"] or {}
    ct = a["calc_totals"] or {}
    g = r.get("gate") or {}
    usage = r.get("batch_usage") or {}
    ws.append([r["boq"], r["sheet"], r["excel_row"], r.get("unit"), g.get("button_reachable"), r.get("run_id"), r.get("run_status"), r.get("run_active"),
               r["description"], " | ".join(note_texts((r.get("payload") or {}).get("notes"))), " > ".join(r["headings"]),
               len(ma), " || ".join(str(x.get("item")) for x in ma), " || ".join(str(x.get("cladding")) for x in ma), " || ".join(str(x.get("thickness_mm")) for x in ma),
               " || ".join(str(x.get("pipe_size_mm")) for x in ma), " || ".join(str(x.get("brand")) for x in ma), " || ".join(str(x.get("material_as_written")) for x in ma),
               a["states"]["item"], a["states"]["cladding"], a["states"]["thickness_mm"], a["states"]["pipe_size_mm"],
               json.dumps(r.get("raw_answer"), ensure_ascii=False)[:1500], json.dumps(a["text_flags"], ensure_ascii=False) if a["text_flags"] else "",
               a["so_verdict"], " || ".join(f"{f.get('attr')}: {f.get('reason')}" for f in a["so_flags"]), json.dumps(r.get("so_failed")) if r.get("so_failed") else "",
               " || ".join(str(b.get("family")) + (" [default]" if b.get("familyDefaulted") else "") for b in blocks), " || ".join(fld(i, "cladding") for i in range(len(blocks))),
               " || ".join(fld(i, "thickness_mm") for i in range(len(blocks))), " || ".join(fld(i, "pipe_size_mm") for i in range(len(blocks))),
               "\n".join(a["rules_fired"]), " || ".join(str(s) for s in a["skus"]), "\n".join("\n".join(w or []) for w in a["working"])[:3000],
               t.get("supply_rate"), t.get("install_rate"), t.get("combined_rate"), a["panel_reason"] if not a["panel_priced"] else "",
               ct.get("supply_rate"), ct.get("install_rate"), ct.get("combined_rate"), a["calc_reason"] if not a["calc_priced"] else "",
               " || ".join(f"{d['what']}: P[{d['panel'][:200]}] C[{d['calculator'][:200]}]" for d in a["divergences"]), a["divergence_cause"] or "",
               a["hand"].get("supply"), a["hand"].get("install"), a["hand_diff"].get("supply"), a["hand_diff"].get("install"),
               a["checks"]["S3_own_thickness"], a["checks"]["material_own"], a["checks"]["cladding"], a["checks"]["pipe_size"], "\n".join(a["violations"]),
               "\n".join(a.get("repeat_diffs", [])) if "repeat_diffs" in a else "(not in the 12d-2 sample)",
               a["schedule_heading"], a["several_values"], a["double_layer_text"], a["foil_mentioned"], r.get("batch_ts"), r.get("batch_rows"),
               f"{usage.get('input_tokens')}/{usage.get('output_tokens')}" if usage else "", r.get("failed_attempts"), sel.get(r["id"], ""),
               HAND["rows"].get(r["id"], {}).get("item", ""), HAND["rows"].get(r["id"], {}).get("cladding", ""),
               HAND["rows"].get(r["id"], {}).get("thickness", ""), HAND["rows"].get(r["id"], {}).get("pipe_size", ""),
               HAND["rows"].get(r["id"], {}).get("note", ""),
               HAND["so"].get(r["id"], {}).get("verdict", ""), HAND["so"].get(r["id"], {}).get("note", "")])
    row = ws.max_row
    if a["violations"] or a["reading_failures"]:
        ws.cell(row=row, column=1).fill = red
    elif a["divergences"] or a["so_verdict"] == "disagree":
        ws.cell(row=row, column=1).fill = amber
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions
for i, c in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = 14 if i < 9 else 28
wb.save(f"{out_dir}/2026-10-10_12d3_Audit_Rows.xlsx")
print("analysis written:", total, "rows; priced", priced, "; reading failures", len(reading_fail), "; violations", len(violations), "; diverged", len(diverged), "; SO disagree", len(so_dis), "; hand-review selection", len(sel))
