# -*- coding: utf-8 -*-
"""PROVENANCE ONLY -- NEVER RUN. Kept for its `attrs_for` mapping table, nothing else.

`nirmaan_stack/services/boq_rate_master/spec_reader.py` states that its deterministic rule set was
carried over VERBATIM from the `attrs_for` table below. That is the only reason this file survives the
2026-09-29 instrument cleanup: it is the provenance record for a shipped module's rules, so a future
maintainer can check spec_reader against what it was derived from.

It is NOT an instrument and must not be executed. It mints `rate_master_hvac_all_v1.json`, which is long
superseded (the live HVAC asset is v14), and it reads a slice-1b copy of the owner's workbook that no
longer needs to exist. Running it would at best rebuild a dead asset.

The original text of the slice-1b build tool follows, unchanged.

Mint rate_master_hvac_all_v1.json (NEW, ADP items) -- the FIRST version of the HVAC series, versioned on
its own (owner amendment 2026-09-21: each discipline keeps its OWN version number; only the file that changed
is minted; NO Electrical file is created or touched). Slice 1b. TEMPORARY, UNTRACKED build tool, re-runnable,
asserts every precondition. Reads the owner's ADP tab from the slice-1b copy of the workbook.

Owner rulings applied (quoted in the slice prompt): R-a separate HVAC file (its own version number, per the
amendment); R-b cost + markup; R-c per-SKU markups, four rate keys; R-d (1) '9/10 NM' -> torque 10,
(2) '1:10 /12' -> panel ratio 12, (3) canvas row 80 -> THREE items SQM/RMT/NOS, (4) eyeball + jet = ONE
family, (6) row 33 '1200MM X300MM' is a FACE size, neck NA, (9) fire-damper TDF flange / fusible link not
modelled; R-e row 34 = 5300/1000, row 38 = 3500/700, row 94 = 6000/800; R-f derived rows 81-86, 89, 91
carry markups ONLY (their sheet formulas are recorded, not applied).

Non-applicable attributes are OMITTED (the db_switchgear_item convention: 30 of 136 items carry only the
keys that apply). The other Electrical convention -- the literal string "NA" on `curve` / `colour` -- is
NOT used because several ADP attributes are numeric and a string in a numeric column would refuse the
CSV round trip (`csv_importer.coerce_attribute` floats every numeric-typed cell).
"""
import hashlib, json, math, os, re, shutil, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "nirmaan_stack", "services", "boq_rate_master", "data")
SRC_XLSX = os.environ.get("HVAC_XLSX") or os.path.join(
    os.environ["LOCALAPPDATA"], "Temp", "claude",
    "C--Users-nites-Documents-frappe-docker-development-frappe-bench-apps-nirmaan-stack",
    "b7734fea-d091-48f5-9a1c-04268035d5bd", "scratchpad", "hvac_pricing_copy_1b.xlsx")
NEW = "v1"
H_NEW = os.path.join(DATA, f"rate_master_hvac_all_{NEW}.json")
KIND = "hvac_adp_item"
DERIVED_ROWS = {81, 82, 83, 84, 85, 86, 89, 91}          # R-f
NEW_COSTS = {34: (5300, 1000), 38: (3500, 700), 94: (6000, 800)}   # R-e
FAMILIES = ["VCD", "fire damper", "actuator", "control panel", "linear grille", "curved grille", "intake louvre",
            "door grille", "floor grille", "slot diffuser", "square diffuser", "round diffuser", "jet / eyeball diffuser",
            "butterfly damper", "spigot", "flexible duct", "disc valve", "NRD", "collar damper", "sound attenuator",
            "cross-talk", "double-skin plenum", "mixing box / LP plenum", "canvas connection", "access door"]

assert not os.path.exists(H_NEW) or "--force" in sys.argv, f"{H_NEW} exists (re-run with --force)"
assert os.path.exists(SRC_XLSX), SRC_XLSX

# ---- read the sheet (values + formulas) ---------------------------------------------------------
wv = openpyxl.load_workbook(SRC_XLSX, data_only=True, read_only=True)["ADP"]
wf = openpyxl.load_workbook(SRC_XLSX, data_only=False, read_only=True)["ADP"]
vals = list(wv.iter_rows(min_row=1, max_row=130, max_col=9, values_only=True))
fmls = list(wf.iter_rows(min_row=1, max_row=130, max_col=9, values_only=True))
hdr = [str(h).strip() if h else "" for h in vals[0]]
assert hdr == ["Item", "ITEM DETAIL", "UNIT", "Supply Markup", "Installation Markup", "BCS Supply",
               "BCS Installation", "BOQ Supply", "BOQ Installation"], hdr
rows = []
for i in range(1, len(vals)):
    v = vals[i]
    if not any(x not in (None, "") for x in v):
        continue
    rec = {"xl_row": i + 1, **{hdr[j]: v[j] for j in range(9)}}
    rec["formulas"] = {hdr[j]: fmls[i][j] for j in range(9) if isinstance(fmls[i][j], str) and fmls[i][j].startswith("=")}
    rows.append(rec)
assert len(rows) == 93, len(rows)

def dia(d):   m = re.search(r"(\d+)\s*mm\s*dia", d, re.I); return int(m.group(1)) if m else None
def wxh(d):   m = re.search(r"(\d+)\s*(?:mm)?\s*x\s*(\d+)", d, re.I); return (int(m.group(1)), int(m.group(2))) if m else None
def neck(d):  m = re.search(r"NECK:\s*(\d+)X(\d+)", d, re.I); return (int(m.group(1)), int(m.group(2))) if m else None
def outer(d): m = re.search(r"OUTER:\s*(\d+)X(\d+)", d, re.I); return (int(m.group(1)), int(m.group(2))) if m else None

def attrs_for(r):
    item = (r["Item"] or "").strip(); det = (r["ITEM DETAIL"] or "").strip(); il, dl = item.lower(), det.lower()
    a = {}
    def fam(f): a["family"] = f
    if il.startswith("volume control damper"):
        fam("VCD"); a["variant"] = {"gi rectangular": "GI rectangular", "gi oval": "GI oval", "motorized": "motorised"}[dl]
    elif il == "fire damper":
        fam("fire damper")
        if "motoris" in dl: a["variant"] = "motorised"; a["ul"] = "no"
        elif "ul 555" in dl: a["variant"] = "UL"; a["ul"] = "yes"
        elif dl.startswith("with ") and "sleeve" in dl: a["variant"] = "with sleeve"; a["ul"] = "no"
        elif "wihtout sleeve" in dl or "without sleeve" in dl: a["variant"] = "without sleeve"; a["ul"] = "no"
        else: raise AssertionError(det)
    elif il.startswith("fire damper actuator"):
        fam("actuator"); a["ul"] = "yes" if "ul 555" in il else "no"
        if dl.startswith("9/10"): a["torque_nm"] = 10.0                           # R-d(1)
        else: a["torque_nm"] = float(re.match(r"([\d.]+)\s*nm", dl).group(1))
    elif "control panel" in il:
        fam("control panel")
        if det.startswith("1:10"): a["panel_ratio"] = 12.0                        # R-d(2)
        else: a["panel_ratio"] = float(re.match(r"1\s*:\s*(\d+)", det).group(1))
    elif il == "grill":
        if "linear" in dl: fam("linear grille"); a["damper"] = "with" if "with damper" in dl else "without"
        elif "curved" in dl: fam("curved grille"); a["damper"] = "with" if ("with damper" in dl and "without" not in dl) else "without"
        elif "intake louver" in dl: fam("intake louvre")
        elif "door grill" in dl: fam("door grille")
        else: raise AssertionError(det)
    elif "slot diffuser" in il:
        fam("slot diffuser"); a["damper"] = "without" if "without" in il else "with"
        m = re.match(r"(\d)\s*slot", dl)
        if m: a["slot_count"] = float(m.group(1))
    elif il.startswith("diffuser with") or il.startswith("diffuser without"):
        fam("square diffuser"); a["damper"] = "without" if "without" in il else "with"
        n, o = neck(det), outer(det)
        if n: assert n[0] == n[1], det; a["neck_mm"] = float(n[0])
        if o: a["face_w_mm"], a["face_h_mm"] = float(o[0]), float(o[1])
        if not n and not o:
            w = wxh(det)
            if w: a["face_w_mm"], a["face_h_mm"] = float(w[0]), float(w[1])   # R-d(6): face, neck NA
    elif il.startswith("round diffuser"):
        fam("round diffuser"); a["damper"] = "without" if "without" in il else "with"; a["dia_mm"] = float(dia(det))
    elif il.startswith("butterfly damper"):
        fam("butterfly damper"); a["dia_mm"] = float(dia(det))
    elif "flexible duct" in il:
        fam("flexible duct"); a["insulated"] = "without" if il.startswith("un-") else "with"; a["dia_mm"] = float(dia(det))
    elif "disc valve" in il:
        fam("disc valve"); a["dia_mm"] = float(dia(det))
    elif "back draft" in il:
        fam("NRD")
    elif il.startswith("collar damper"):
        fam("collar damper")
    elif il == "sound attenuator":
        fam("sound attenuator")
    elif il.startswith("double skin plenum"):
        fam("double-skin plenum"); a["thickness_mm"] = float(re.match(r"(\d+)\s*mm", dl).group(1))
    elif il.startswith("ms floor grill"):
        fam("floor grille"); a["damper"] = "without" if "without" in dl else "with"
    elif il == "spigot":
        fam("spigot"); a["dia_mm"] = float(dia(det))
    elif "canvas" in il:
        fam("canvas connection")
    elif il.startswith("z-piece"):
        fam("cross-talk"); w = wxh(det)
        if w: a["face_w_mm"], a["face_h_mm"] = float(w[0]), float(w[1])
    elif il.startswith("low pressure plenum"):
        fam("mixing box / LP plenum"); a["insulated"] = "without" if "without" in il else "with"
        m = re.search(r"(\d+)\s*X\s*(\d+)\s*X\s*(\d+)", det, re.I)
        if m: a["face_w_mm"], a["face_h_mm"], a["depth_mm"] = float(m.group(1)), float(m.group(2)), float(m.group(3))
    elif "eye ball" in il:
        fam("jet / eyeball diffuser"); a["dia_mm"] = float(dia(det))          # R-d(4)
    elif "access door" in il:
        fam("access door"); w = wxh(det); a["face_w_mm"], a["face_h_mm"] = float(w[0]), float(w[1])
    else:
        raise AssertionError(f"no family: {item!r} / {det!r}")
    assert a["family"] in FAMILIES, a["family"]
    return a

def uid(kind, row, unit):
    return "rmi-" + hashlib.sha1(f"{kind}|ADP|{row}|{unit}".encode()).hexdigest()[:12]

def roundup0(x): return math.ceil(x - 1e-9)

items, expected, formulas, derived = [], {}, {}, []
for r in rows:
    xl = r["xl_row"]; ms, mi = r["Supply Markup"], r["Installation Markup"]
    assert isinstance(ms, (int, float)) and isinstance(mi, (int, float)), xl
    cs, ci = r["BCS Supply"], r["BCS Installation"]
    if xl in NEW_COSTS: cs, ci = NEW_COSTS[xl]
    units = ["SQM", "RMT", "NOS"] if xl == 80 else [str(r["UNIT"]).strip()]            # R-d(3)
    a = attrs_for(r)
    for u in units:
        rates = {"supply_markup": float(ms), "install_markup": float(mi)}
        if xl in DERIVED_ROWS:
            derived.append(xl)
            formulas[xl] = {k: v for k, v in r["formulas"].items() if k in ("BCS Supply", "BCS Installation")}
        else:
            assert isinstance(cs, (int, float)) and isinstance(ci, (int, float)), (xl, cs, ci)
            rates["cost_supply"], rates["cost_install"] = float(cs), float(ci)
            expected[(xl, u)] = (roundup0(cs * (1 + ms)), roundup0(ci * (1 + mi)))
        items.append({"kind": KIND, "brand": None, "unit": u, "attributes": dict(a), "rates": rates,
                      "source": {"sheet": "ADP", "row": xl}, "item_uid": uid(KIND, xl, u)})
assert len(items) == 95, len(items)
assert len({i["item_uid"] for i in items}) == 95
# the exporter's order (`kind asc, item_uid asc`), so a re-export of the loaded catalogue reproduces this file
items.sort(key=lambda i: (i["kind"], i["item_uid"]))
# the sheet's own BOQ figures must be reproduced (rows with a numeric BOQ cell; row 9's install cell is corrupt text)
sheet_boq = {r["xl_row"]: (r["BOQ Supply"], r["BOQ Installation"]) for r in rows}
mismatch = []
for (xl, u), (es, ei) in expected.items():
    if xl in NEW_COSTS: continue
    bs, bi = sheet_boq[xl]
    if isinstance(bs, (int, float)) and bs != es: mismatch.append((xl, "supply", bs, es))
    if isinstance(bi, (int, float)) and bi != ei: mismatch.append((xl, "install", bi, ei))
assert not mismatch, mismatch
for xl, (es, ei) in {34: (7685, 1600), 38: (5075, 1120), 94: (8700, 1280)}.items():
    assert expected[(xl, str(next(r for r in rows if r["xl_row"] == xl)["UNIT"]).strip())] == (es, ei), xl

config = {
    "category_id": "hvac_adp",
    "category_display": "ADP (Air Distribution Products)",
    "item_kinds": [KIND],
    "attribute_definitions": [
        {"id": "family", "label": "Family", "type": "choice", "values": FAMILIES},
        {"id": "damper", "label": "Damper", "type": "choice", "values": ["with", "without"]},
        {"id": "insulated", "label": "Insulated", "type": "choice", "values": ["with", "without"]},
        {"id": "neck_mm", "label": "Neck (mm)", "type": "number"},
        {"id": "face_w_mm", "label": "Face / size width (mm)", "type": "number"},
        {"id": "face_h_mm", "label": "Face / size height (mm)", "type": "number"},
        {"id": "depth_mm", "label": "Depth (mm)", "type": "number"},
        {"id": "dia_mm", "label": "Diameter (mm)", "type": "number"},
        {"id": "slot_count", "label": "Slot count", "type": "number"},
        {"id": "torque_nm", "label": "Torque (NM)", "type": "number"},
        {"id": "ul", "label": "UL 555 rated", "type": "choice", "values": ["yes", "no"]},
        {"id": "panel_ratio", "label": "Panel ratio (1:N)", "type": "number"},
        {"id": "thickness_mm", "label": "Thickness (mm)", "type": "number"},
        {"id": "variant", "label": "Variant", "type": "choice",
         "values": ["GI rectangular", "GI oval", "motorised", "with sleeve", "without sleeve", "UL"]},
    ],
    "pipelines": {},
    "notes": ("Slice 1b (2026-09-21): DATA-ONLY -- items + attributes, no pipelines, so the config is NOT eligible "
              "for pricing or extraction (extraction.config_is_eligible / pricingSheetHelper.isEligibleConfig both "
              "require non-empty pipelines). Rates per SKU: cost_supply, cost_install, supply_markup, install_markup "
              "(R-c); the eight derived items (sheet rows 81-86, 89, 91) carry markups only (R-f). Pricing rule for "
              "a later slice: BoQ = ROUNDUP(cost x (1 + markup), 0), per side."),
    "discipline": "HVAC",
}
attr_ids = {d["id"] for d in config["attribute_definitions"]}
rate_keys = {k for i in items for k in i["rates"]}
assert not (attr_ids & rate_keys), attr_ids & rate_keys
for i in items:
    assert set(i["attributes"]) <= attr_ids, i
hvac = {
    "discipline": "HVAC",
    "source_workbook": "HVAC_BOQ_BCS PRICING_ Nitesh EditsV2.xlsx",
    "items": items,
    "category_configs": [config],
    "goldens": {},
    "retired_kinds": [],
    "retired_category_ids": [],
    "retirement_reasons": {"kinds": {}, "categories": {}},
}
with open(H_NEW, "w", encoding="utf-8", newline="\n") as fh:
    json.dump(hvac, fh, indent=1, ensure_ascii=False); fh.write("\n")

fam_counts = {}
for i in items: fam_counts[i["attributes"]["family"]] = fam_counts.get(i["attributes"]["family"], 0) + 1
report = {
    "items": len(items), "derived_items": len(derived), "priced_items": len(expected),
    "families": fam_counts, "rate_keys": sorted(rate_keys), "attr_ids": sorted(attr_ids),
    "derived_formulas": formulas,
    "expected_boq": {f"{xl}|{u}": v for (xl, u), v in sorted(expected.items())},
    "sha256_hvac": hashlib.sha256(open(H_NEW, "rb").read()).hexdigest(),
}
with open(os.path.join(os.path.dirname(SRC_XLSX), "mint_hvac_v1_report.json"), "w", encoding="utf-8") as fh:
    json.dump(report, fh, indent=1, ensure_ascii=False)
print(json.dumps({k: v for k, v in report.items() if k not in ("expected_boq",)}, indent=1, ensure_ascii=False))
